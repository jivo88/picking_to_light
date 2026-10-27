"""Gera Orcamento_Pecas_Fontes.xlsx: orçamento das peças (Arquitetura A + RFID futuro) com o link de cada preço.

Preços: fontes_precos.py (consulta de 27/10/2026). Quantidades: lidas do orçamento completo já recalculado
(Orcamento_PTL.xlsx, abas A-Interno e RFID-futuro) e da composição de cada módulo.
Uso: python3 gerar_pecas_fontes.py [completo.xlsx] [saida.xlsx]
"""
import os
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import fontes_precos as FP  # noqa: E402

BASE = Path(__file__).resolve().parent.parent
ORIG = Path(sys.argv[1]) if len(sys.argv) > 1 else BASE / "entregaveis" / "Orcamento_PTL.xlsx"
SAIDA = Path(sys.argv[2]) if len(sys.argv) > 2 else BASE / "entregaveis" / "Orcamento_Pecas_Fontes.xlsx"

src = openpyxl.load_workbook(ORIG, data_only=True)
ai, rf, pr = src["A-Interno"], src["RFID-futuro"], src["Premissas"]
if ai["G21"].value is None:
    sys.exit("Orcamento_PTL.xlsx sem valores calculados: rode recalc.py nele antes.")


def premissa(inicio):
    for row in pr.iter_rows(min_row=2, values_only=True):
        if row[0] and str(row[0]).startswith(inicio):
            return row[1], row[3]
    raise KeyError(inicio)


F = "Arial"
AZUL = Font(name=F, size=10, color="0000FF")
PRETO = Font(name=F, size=10)
VERDE = Font(name=F, size=10, color="008000")
NEG = Font(name=F, size=10, bold=True)
TIT = Font(name=F, size=14, bold=True, color="1F3864")
CAB = Font(name=F, size=10, bold=True, color="FFFFFF")
NOTA = Font(name=F, size=9, italic=True, color="595959")
LINK = Font(name=F, size=10, color="0563C1", underline="single")
F_CAB = PatternFill("solid", fgColor="1F3864")
F_SEC = PatternFill("solid", fgColor="D9E1F2")
F_TOT = PatternFill("solid", fgColor="FCE4D6")
F_EST = PatternFill("solid", fgColor="FFF2CC")
FINA = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
BRL = 'R$ #,##0.00;(R$ #,##0.00);"-"'
NUM = '#,##0.##;(#,##0.##);"-"'
PCT = '0.0%;(0.0%);"-"'


def larguras(ws, ls):
    for i, w in enumerate(ls, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, ref, v, fonte=PRETO, fmt=None, fill=None, borda=True, wrap=False):
    c = ws[ref]
    c.value = v
    c.font = fonte
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if borda:
        c.border = BORDA
    c.alignment = Alignment(wrap_text=wrap, vertical="top")
    return c


def cabecalho(ws, r, tit, altura=32):
    for i, t in enumerate(tit, 1):
        c = ws.cell(row=r, column=i, value=t)
        c.font, c.fill, c.border = CAB, F_CAB, BORDA
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = altura


def faixa(ws, r, ncol, rotulo, fill):
    for i in range(1, ncol + 1):
        ws.cell(row=r, column=i).fill = fill
        ws.cell(row=r, column=i).border = BORDA
    ws.cell(row=r, column=1, value=rotulo).font = NEG


def nota(ws, r, texto, ate, altura=28):
    put(ws, f"A{r}", texto, NOTA, borda=False, wrap=True)
    ws.merge_cells(f"A{r}:{ate}{r}")
    ws.row_dimensions[r].height = altura


# ---------------------------------------------------------------- quantidades por código
def q(row):          # (piloto, completo) de uma linha da A-Interno
    return ai.cell(row=row, column=3).value or 0, ai.cell(row=row, column=4).value or 0


MOD = {"q": q(3), "no": q(4), "eq": q(5), "conc": q(6)}
BOM = {
    "q": [("P01", 1), ("P09", 1), ("P13", 1), ("P10", 1), ("P12", 2), ("P11", 2)],
    "no": [("P03", 1), ("P04", 1), ("P08", 1), ("P15", 1), ("P10", 6), ("P14", 1)],
    "eq": [("P02", 1), ("P01", 1), ("P03", 1), ("P04", 1), ("P08", 1), ("P09", 1), ("P10", 2), ("P14", 1), ("P15", 1)],
    "conc": [("P05", 1), ("P04", 2), ("P06", 1), ("P08", 1), ("P16", 1), ("P17", 1), ("P10", 2), ("P38", 3), ("P39", 1),
             ("P40", 1)],
}
DIRETOS = [("P12", 7), ("P11", 8), ("P18", 9), ("P19", 10), ("P47", 11), ("P12", 12), ("P20", 13), ("P43", 14),
           ("P44", 15), ("P12", 16), ("P46", 17), ("P41", 18), ("P42", 19)]
BANCADA_EXTRA = {"P47": 1, "P12": 20, "P11": 20, "P48": 1, "P49": 1}   # fase 1 da POC (módulos reaproveitados não contam)
USO = {"P01": "Módulo Q, Módulo E+Q", "P02": "Módulo E+Q", "P03": "Nó de coluna, Módulo E+Q",
       "P04": "Nó de coluna, Módulo E+Q, Concentrador (2)", "P05": "Concentrador", "P06": "Concentrador",
       "P08": "Nó de coluna, Módulo E+Q, Concentrador", "P09": "Módulo Q, Módulo E+Q",
       "P10": "Módulo Q (1), Nó (6), E+Q (2), Concentrador (2)", "P11": "barramento e bancada",
       "P12": "barramento, ligações dos módulos Q, rede até a estação, uplink, bancada", "P13": "Módulo Q",
       "P14": "Nó de coluna, Módulo E+Q", "P15": "Nó de coluna, Módulo E+Q", "P16": "Concentrador",
       "P17": "Concentrador", "P18": "fixação dos módulos", "P19": "1 kit por estação", "P20": "rede das estações",
       "P38": "alimentação do concentrador (3 m)", "P39": "alimentação do concentrador",
       "P40": "alimentação do concentrador", "P41": "ferramenta (compra única)", "P42": "ferramenta (compra única)",
       "P43": "1 por switch", "P44": "1 por switch", "P46": "estações sem tomada próxima",
       "P47": "1 por estação", "P48": "bancada da POC", "P49": "bancada da POC",
       "P22": "1 por caixa da frota", "P23": "1 por caixa da frota", "P24": "pulseira", "P25": "pulseira",
       "P26": "pulseira", "P27": "pulseira", "P45": "1 por estação"}

QTD = {}
for mod, itens in BOM.items():
    for cod, k in itens:
        a, b = QTD.get(cod, (0, 0))
        QTD[cod] = (a + MOD[mod][0] * k, b + MOD[mod][1] * k)
for cod, row in DIRETOS:
    a, b = QTD.get(cod, (0, 0))
    pa, pb = q(row)
    QTD[cod] = (a + pa, b + pb)
for cod, n in BANCADA_EXTRA.items():
    a, b = QTD.get(cod, (0, 0))
    QTD[cod] = (a + n, b)

qr = lambda r: (rf.cell(row=r, column=3).value or 0, rf.cell(row=r, column=4).value or 0)  # noqa: E731
TAGS, PULS, CARR = qr(3), qr(5), qr(6)
QTD_RF = {"P22": TAGS, "P23": TAGS, "P24": PULS, "P25": PULS, "P26": PULS, "P27": PULS, "P45": CARR}

ORDEM_A = ["P01", "P02", "P03", "P04", "P05", "P06", "P08", "P09", "P10", "P11", "P12", "P13", "P14", "P15", "P16",
           "P17", "P18", "P19", "P20", "P38", "P39", "P40", "P41", "P42", "P43", "P44", "P46", "P47", "P48", "P49"]
ORDEM_RF = ["P22", "P23", "P24", "P25", "P26", "P27", "P45"]

wb = openpyxl.Workbook()

# ================================================================ Premissas
wp = wb.active
wp.title = "Premissas"
larguras(wp, [44, 14, 90])
put(wp, "A1", "Premissas usadas nos preços", TIT, borda=False)
cabecalho(wp, 3, ["Premissa", "Valor", "Fonte"])
cambio, f_c = premissa("Câmbio")
fator, f_f = premissa("Fator de importação")
frete, f_fr = premissa("Frete das compras")
PREM = [("Câmbio USD/BRL", cambio, '"R$ "0.0000', f_c), ("Fator de importação sobre o valor em US$", fator, '0.00"x"', f_f),
        ("Frete das compras nacionais", frete, PCT, f_fr)]
for i, (rot, v, fmt, fonte) in enumerate(PREM):
    put(wp, f"A{4 + i}", rot)
    put(wp, f"B{4 + i}", v, AZUL, fmt)
    put(wp, f"C{4 + i}", fonte, NOTA, wrap=True)
CAMBIO, FATOR, FRETE = "Premissas!$B$4", "Premissas!$B$5", "Premissas!$B$6"
nota(wp, 8, "Só a placa W5500-EVB-Pico (P05) é importada; as outras peças têm preço em reais em lojas do Brasil.", "C")

# ================================================================ Fontes
wf = wb.create_sheet("Fontes")
larguras(wf, [8, 30, 16, 52, 44, 14, 18, 12, 12, 14, 11, 62])
cabecalho(wf, 1, ["Código", "Componente", "Loja", "Produto (como aparece na loja)", "Link", "Preço publicado",
                  "Forma do preço", "Unidades na embalagem", "Qtd por unidade do código", "Preço no código (R$)",
                  "Data", "Observação"], 42)
for i, c in enumerate(FP.C, 2):
    est = c["forma"].startswith("estimativa")
    put(wf, f"A{i}", c["cod"])
    put(wf, f"B{i}", c["componente"], wrap=True)
    put(wf, f"C{i}", c["loja"])
    put(wf, f"D{i}", c["produto"], wrap=True)
    lk = put(wf, f"E{i}", c["url"] or "sem link (cotar com a manutenção predial)", wrap=True)
    if c["url"]:
        lk.hyperlink = c["url"]
        lk.font = LINK
    put(wf, f"F{i}", c["preco"], AZUL, '"US$ "#,##0.00' if c["moeda"] == "USD" else BRL, F_EST if est else None)
    put(wf, f"G{i}", c["forma"], fill=F_EST if est else None, wrap=True)
    put(wf, f"H{i}", c["divisor"], AZUL, NUM)
    put(wf, f"I{i}", c["qtd"], AZUL, NUM)
    conv = f"*{CAMBIO}*{FATOR}" if c["moeda"] == "USD" else ""
    put(wf, f"J{i}", f"=F{i}/H{i}*I{i}{conv}", PRETO, BRL)
    put(wf, f"K{i}", FP.DATA)
    put(wf, f"L{i}", c["obs"], NOTA, wrap=True)
FF = len(FP.C) + 1
wf.freeze_panes = "C2"
wf.auto_filter.ref = f"A1:L{FF}"

# ================================================================ Peças
wq = wb.create_sheet("Peças", 0)
larguras(wq, [8, 50, 8, 14, 12, 12, 16, 16, 34, 26])
put(wq, "A1", "Orçamento das peças com fontes: Arquitetura A (equipe interna) e RFID futuro", TIT, borda=False)
nota(wq, 2, f"Preços consultados em {FP.DATA}; o preço de cada código soma os componentes da aba Fontes, onde estão a loja, "
            "o produto e o link de cada um. Quantidades do orçamento completo (com 10% de reserva). Só peças e materiais: "
            "mão de obra e software ficam fora.", "J", 40)
cabecalho(wq, 4, ["Código", "Peça", "Unid.", "Preço unit. (R$)", "Qtd POC (04 e 27 + bancada)", "Qtd projeto (40 estações)",
                  "Total POC (R$)", "Total projeto (R$)", "Onde é usada", "Fonte"], 44)


def bloco(r, titulo, ordem, qtds):
    for i in range(1, 11):
        wq.cell(row=r, column=i).fill = F_SEC
        wq.cell(row=r, column=i).border = BORDA
    wq.cell(row=r, column=1, value=titulo).font = NEG
    r += 1
    ini = r
    for cod in ordem:
        desc, un = FP.CODIGOS[cod]
        comps = [c for c in FP.C if c["cod"] == cod]
        est = any(c["forma"].startswith("estimativa") for c in comps)
        put(wq, f"A{r}", cod)
        put(wq, f"B{r}", desc, wrap=True)
        put(wq, f"C{r}", un)
        put(wq, f"D{r}", f"=SUMIF(Fontes!$A$2:$A${FF},A{r},Fontes!$J$2:$J${FF})", VERDE, BRL, F_EST if est else None)
        a, b = qtds.get(cod, (0, 0))
        put(wq, f"E{r}", round(a, 2), AZUL, NUM)
        put(wq, f"F{r}", round(b, 2), AZUL, NUM)
        put(wq, f"G{r}", f"=D{r}*E{r}", PRETO, BRL)
        put(wq, f"H{r}", f"=D{r}*F{r}", PRETO, BRL)
        put(wq, f"I{r}", USO.get(cod, ""), NOTA, wrap=True)
        if len(comps) == 1 and comps[0]["url"]:
            lk = put(wq, f"J{r}", comps[0]["loja"], LINK)
            lk.hyperlink = comps[0]["url"]
        elif not comps[0]["url"]:
            put(wq, f"J{r}", "cotar (sem link)", fill=F_EST)
        else:
            put(wq, f"J{r}", f"{len(comps)} componentes, aba Fontes" + (" (1 estimado)" if est else ""), wrap=True,
                fill=F_EST if est else None)
        r += 1
    fim = r - 1
    put(wq, f"B{r}", "Frete das compras (% da aba Premissas)")
    for col in "ACDEFIJ":
        put(wq, f"{col}{r}", None)
    put(wq, f"G{r}", f"=SUM(G{ini}:G{fim})*{FRETE}", PRETO, BRL)
    put(wq, f"H{r}", f"=SUM(H{ini}:H{fim})*{FRETE}", PRETO, BRL)
    r += 1
    faixa(wq, r, 10, "Total " + titulo.lower(), F_TOT)
    put(wq, f"G{r}", f"=SUM(G{ini}:G{r - 1})", NEG, BRL, F_TOT)
    put(wq, f"H{r}", f"=SUM(H{ini}:H{r - 1})", NEG, BRL, F_TOT)
    return r


R_A = bloco(5, "Peças e materiais da Arquitetura A", ORDEM_A, QTD)
R_RF = bloco(R_A + 2, "RFID futuro (fora do investimento inicial)", ORDEM_RF, QTD_RF)
nota(wq, R_RF + 2, "Colunas da POC: estações 04 e 27 com reserva, mais os itens só da bancada (sinal, cabo, plugs, painel de MDF "
                   "e switch de 5 portas); o concentrador, o nó e os módulos da bancada vão para a estação 04 e não são "
                   "contados duas vezes. No RFID, a coluna POC é o piloto do RFID nas mesmas 2 estações.", "J", 40)
nota(wq, R_RF + 3, "Fundo amarelo: valor sem preço lido em loja. P46 é serviço de eletricista (cotar com a manutenção "
                   "predial). A bateria LiPo do P26 tem link, mas o site não respondeu na consulta.", "J", 28)
wq.freeze_panes = "C5"

# ================================================================ Resumo
wr = wb.create_sheet("Resumo", 0)
larguras(wr, [52, 18, 18, 60])
put(wr, "A1", "Peças do Pick-to-Light com fontes", TIT, borda=False)
nota(wr, 2, f"Preços de lojas on-line consultados em {FP.DATA}. Cada componente tem o link na aba Fontes.", "D")
cabecalho(wr, 4, ["Item", "POC (R$)", "Projeto 40 estações (R$)", "Observação"])
LR = [("Peças e materiais da Arquitetura A (com frete)", f"='Peças'!G{R_A}", f"='Peças'!H{R_A}",
       "igual à linha de materiais do orçamento completo e do projeto interno"),
      ("RFID futuro: peças (com frete)", f"='Peças'!G{R_RF}", f"='Peças'!H{R_RF}", "fora do investimento inicial")]
for i, (rot, a, b, obs) in enumerate(LR):
    put(wr, f"A{5 + i}", rot)
    put(wr, f"B{5 + i}", a, VERDE, BRL)
    put(wr, f"C{5 + i}", b, VERDE, BRL)
    put(wr, f"D{5 + i}", obs, NOTA, wrap=True)
n_comp = len(FP.C)
n_link = sum(1 for c in FP.C if c["url"])
n_lido = sum(1 for c in FP.C if c["url"] and not c["forma"].startswith("estimativa"))
put(wr, "A8", "Componentes pesquisados", NEG, borda=False)
for i, (rot, v) in enumerate([("Componentes", n_comp), ("Com link", n_link), ("Com link e preço lido na página", n_lido),
                              ("Sem link (serviço a cotar)", n_comp - n_link),
                              ("Com link, preço não lido (estimado)", n_link - n_lido)]):
    put(wr, f"A{9 + i}", rot)
    put(wr, f"B{9 + i}", v, PRETO, '0')
nota(wr, 15, "Os preços são de varejo, para 1 unidade, sem frete (o frete entra como percentual). Em compras de centenas "
             "de peças vale pedir cotação de atacado às mesmas lojas. Vários itens estavam com 'últimas unidades' na "
             "consulta: ver a coluna Observação da aba Fontes.", "D", 42)
nota(wr, 16, "Abas: Peças (quantidades e totais) · Fontes (loja, produto, link e preço de cada componente) · Premissas.", "D")

for w in wb.worksheets:
    w.sheet_view.showGridLines = False
    w.page_setup.orientation = "landscape"
    w.page_setup.fitToWidth, w.page_setup.fitToHeight = 1, 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(SAIDA)
print("ok", SAIDA, "componentes", n_comp, "com link", n_link, "lidos", n_lido)
