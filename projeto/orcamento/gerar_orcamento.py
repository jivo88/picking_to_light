"""Gera Orcamento_PTL.xlsx com formulas (precos, composicao, cenarios). Revisão 26/10/2026."""
import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

EST = json.load(open("/home/claude/ptl-projeto/dados/estacoes.json", encoding="utf-8"))
SAIDA = "/home/claude/ptl-projeto/entregaveis/Orcamento_PTL.xlsx"

AZUL = Font(name="Arial", size=10, color="0000FF")
PRETO = Font(name="Arial", size=10)
NEG = Font(name="Arial", size=10, bold=True)
VERDE = Font(name="Arial", size=10, color="008000")
TIT = Font(name="Arial", size=14, bold=True, color="1F3864")
SUB = Font(name="Arial", size=11, bold=True, color="1F3864")
CAB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
NOTA = Font(name="Arial", size=9, italic=True)
FUNDO_CAB = PatternFill("solid", fgColor="1F3864")
AMARELO = PatternFill("solid", fgColor="FFF2CC")
TOTAL = PatternFill("solid", fgColor="DDEBF7")
FINA = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)
BRL = 'R$ #,##0.00;(R$ #,##0.00);"-"'
BRL0 = 'R$ #,##0;(R$ #,##0);"-"'
NUM = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%'

wb = Workbook()


def cabecalho(ws, linha, titulos, larguras=None):
    for i, t in enumerate(titulos, 1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font = CAB
        c.fill = FUNDO_CAB
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA
    ws.row_dimensions[linha].height = 32
    if larguras:
        for i, w in enumerate(larguras, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def cel(ws, ref, valor, fonte=PRETO, fmt=None, fill=None, borda=True, wrap=False):
    c = ws[ref]
    c.value = valor
    c.font = fonte
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if borda:
        c.border = BORDA
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    return c


def faixa_total(ws, r, rotulo, cols="BCDE"):
    cel(ws, "A%d" % r, rotulo, NEG, fill=TOTAL)
    for col in cols:
        cel(ws, "%s%d" % (col, r), None, fill=TOTAL)


# =====================================================================
# 1. LEIA-ME
# =====================================================================
ws = wb.active
ws.title = "Leia-me"
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 115
linhas = [
    ("Orçamento do projeto Pick-to-Light (flow rack)", TIT),
    ("Arquitetura A (coletor + concentrador + displays) e Arquitetura B (identificação por RFID). Revisão 26/10/2026.", SUB),
    ("", None),
    ("COMO USAR", SUB),
    ("• Células em AZUL são entradas editáveis (quantidades, preços, premissas). Células em PRETO são fórmulas: não editar.", PRETO),
    ("• Células com fundo AMARELO são premissas/estimativas que precisam de cotação ou confirmação.", PRETO),
    ("• Aba 'Estações': números de endereços por estação, extraídos de wms_endereco (esteiras 1 e 2) em 23/10/2026. "
     "Corrija qualquer número direto ali: todo o orçamento recalcula.", PRETO),
    ("• Coluna 'Piloto' (S/N) na aba 'Estações' define quais estações entram no cenário piloto (POC: 04 e 27).", PRETO),
    ("• Aba 'Preços': catálogo de peças com fonte e data. Itens marcados 'Estimativa' devem ser cotados antes da compra.", PRETO),
    ("• A operação completa inclui as estações do piloto (não somar os dois).", PRETO),
    ("", None),
    ("ABAS", SUB),
    ("Resumo: comparativo de todos os cenários (piloto x operação completa), contingência e investimento total recomendado.", PRETO),
    ("Premissas: câmbio, fator de importação, custos-hora, horas de montagem/instalação, metragens, reserva, frete, contingência.", PRETO),
    ("Estações: dimensionamento por estação (40 estações: 20 Medicamento + 20 Perfumaria).", PRETO),
    ("Preços: catálogo de peças com fonte.", PRETO),
    ("Composição: lista de materiais (BOM) de cada tipo de módulo e custo unitário.", PRETO),
    ("A-Interno: Arquitetura A com produção interna: peças, infraestrutura, mão de obra, treinamento, variante 5 V.", PRETO),
    ("A-Terceiros: Arquitetura A com produção de terceiros: (T1) placa própria montada por EMS e (T2) PTL comercial importado.", PRETO),
    ("B-RFID: Arquitetura B: mesma infraestrutura de displays + identificação RFID (B1 pulseira, B2 leitor fixo).", PRETO),
    ("Software: esforço estimado de desenvolvimento (horas) para as duas arquiteturas.", PRETO),
    ("", None),
    ("TERMOS", SUB),
    ("Módulo Q: display de QUANTIDADE (3 dígitos) + botão com LED, um por endereço de alto giro (andares 0 a 3).", PRETO),
    ("Módulo E+Q: display de ENDEREÇO (8 dígitos) + QUANTIDADE (3 dígitos) + botão; usado nos andares 4 e 5 do alto giro "
     "(2 zonas por estação: P001–P005 e P006–P010, por baixo da esteira de separação) e em cada zona de baixo giro "
     "(2 zonas por lado do corredor, na altura dos olhos).", PRETO),
    ("Nó de coluna: microcontrolador por coluna (prédio) do alto giro que liga até 4 Módulos Q ao barramento RS-485 "
     "(nos dados, nenhum prédio tem mais de 4 endereços nos andares 0-3).", PRETO),
    ("Concentrador: controlador da estação (Ethernet cabeada), recebe a caixa do serviço PTL e comanda os 2 barramentos.", PRETO),
    ("Zona de baixo giro: 2 zonas por lado do corredor (cerca de 6 prédios cada), cada uma com 1 display E+Q na altura dos olhos "
     "(o prédio 015 entra na zona D1; esquinas têm zonas de 4 prédios).", PRETO),
    ("Sinal de baixo giro: luz âmbar ao lado da posição de pega que acende quando a caixa tem itens no corredor.", PRETO),
    ("", None),
    ("O QUE MUDOU NA REVISÃO DE 26/10/2026", SUB),
    ("• Canaleta calculada pelo layout real de cada estação (frente do alto giro + descidas por prédio + 2 lados do corredor + travessia), "
     "com as medidas do desenho de 26/10/2026: estação de 2,5 m, corredor de 1,2 m, prateleiras de 50 cm × 2 m (cerca de 27 m por estação).", PRETO),
    ("• Concentrador atrás da estação: travessia de 6 m por cima das esteiras até o corredor.", PRETO),
    ("• Barramento: 1 cabo por nó/módulo E+Q + cabo de saída das 2 portas do concentrador (antes: +2 segmentos genéricos).", PRETO),
    ("• Itens que faltavam: alimentação do concentrador (cabo PP, plugue, DPS), pontos de tomada novos, nobreak e mini rack "
     "dos switches, cabo de uplink, ferramentas de montagem e frete das compras.", PRETO),
    ("• Botão de confirmação trocado por botão metálico com anel de LED (cerca de 6.000 confirmações por estação por dia).", PRETO),
    ("• Instalação calculada por quantidade de módulos e metros de canaleta, com hora extra (instalação fora do horário de separação).", PRETO),
    ("• Treinamento de operadores e manutenção incluído em todos os cenários.", PRETO),
    ("• B1: 2 pulseiras por estação (a bateria de 500 mAh não cobre um turno inteiro com leitura contínua) + carregador.", PRETO),
    ("• B2: o leitor fixo vira um nó do barramento RS-485 (microcontrolador próprio), sem ligação I2C longa até o concentrador.", PRETO),
    ("• T1: horas de montagem final/gravação dos módulos vindos da EMS; T1/T2: cabo em cadeia dos módulos Q com 1 m por módulo.", PRETO),
    ("• Software: custo-hora calculado com o salário de mercado de analista de desenvolvimento (antes R$ 60/h fixo) "
     "e implantação do servidor/broker incluída.", PRETO),
    ("• Câmbio atualizado (fechamento de 23/10/2026) e contingência de 10% no Resumo.", PRETO),
    ("• Ajuste de 26/10/2026 (tarde): 2 displays E+Q nos andares 4-5 (antes 4) e 2 zonas por lado no baixo giro (antes 4), "
     "o que reduz os módulos E+Q de 485 para 251; sinal luminoso de baixo giro por estação.", PRETO),
]
for i, (t, f) in enumerate(linhas, 1):
    c = ws.cell(row=i + 1, column=2, value=t)
    if f:
        c.font = f
    c.alignment = Alignment(wrap_text=True, vertical="top")

# =====================================================================
# 2. PREMISSAS
# =====================================================================
wp = wb.create_sheet("Premissas")
cabecalho(wp, 1, ["Premissa", "Valor", "Unidade", "Fonte / observação"], [52, 14, 12, 95])
premissas = [
    # chave, descricao, valor (formula com {chave}), unidade, fonte, amarelo?
    ("cambio", "Câmbio USD/BRL (comercial)", 5.0371, "R$/US$",
     "renovainvest.com.br/dolar-hoje, fechamento de 23/10/2026 (13h07, mercado fechado)", False),
    ("fator_imp", "Fator de importação (frete + impostos + despacho) sobre o valor FOB", 1.9, "x",
     "Estimativa para importação formal de eletrônicos (II + IPI + PIS/COFINS + ICMS + frete). Confirmar com Comércio Exterior.", True),
    ("salario_tec", "Salário médio técnico eletrônico (CLT)", 2846.84, "R$/mês",
     "salario.com.br, CBO 313215, atualizado em 08/10/2026 (CAGED)", False),
    ("salario_dev", "Salário médio analista de desenvolvimento de sistemas (CLT)", 8458.30, "R$/mês",
     "salario.com.br, CBO 212405, atualizado em 08/10/2026 (CAGED)", False),
    ("salario_oper", "Salário médio auxiliar de logística (operador da estação)", 2009.59, "R$/mês",
     "salario.com.br, CBO 414140, atualizado em 08/10/2026 (CAGED)", False),
    ("encargos", "Fator de encargos e benefícios sobre salário", 2.0, "x",
     "Estimativa usual (encargos + benefícios + provisões). Ajustar com RH.", True),
    ("horas_mes", "Horas produtivas por mês", 163, "h", "44 h/semana com ausências", True),
    ("custo_hora", "Custo-hora técnico (calculado)", "={salario_tec}*{encargos}/{horas_mes}", "R$/h",
     "Fórmula: salário × encargos ÷ horas", False),
    ("custo_hora_dev", "Custo-hora desenvolvimento interno (calculado)", "={salario_dev}*{encargos}/{horas_mes}", "R$/h",
     "Fórmula: salário × encargos ÷ horas (antes: R$ 60/h fixo)", False),
    ("custo_hora_oper", "Custo-hora operador (calculado)", "={salario_oper}*{encargos}/{horas_mes}", "R$/h",
     "Usado no tempo de treinamento dos operadores", False),
    ("h_mod_q", "Montagem, Módulo Q", 0.25, "h/un", "Estimativa: display + botão + carcaça + teste", True),
    ("h_no", "Montagem, Nó de coluna", 0.75, "h/un", "Estimativa: placa base, soldas, gravação de firmware, teste", True),
    ("h_mod_eq", "Montagem, Módulo E+Q", 1.0, "h/un", "Estimativa", True),
    ("h_conc", "Montagem, Concentrador (gabinete completo)", 2.0, "h/un", "Estimativa", True),
    ("h_inst_est", "Instalação, parte fixa por estação", 4, "h/estação",
     "Concentrador, energia, cabo de rede, passagem da frente para o corredor", True),
    ("h_inst_q", "Instalação por Módulo Q (fixar, ligar, etiquetar)", 0.15, "h/un", "Estimativa", True),
    ("h_inst_no", "Instalação por Nó de coluna", 0.25, "h/un", "Estimativa", True),
    ("h_inst_eq", "Instalação por Módulo E+Q", 0.35, "h/un", "Estimativa", True),
    ("h_inst_m", "Instalação de canaleta", 0.08, "h/m", "Estimativa (cerca de 5 min por metro)", True),
    ("h_comiss", "Testes e comissionamento por estação", 4, "h/estação", "Teste de cada endereço + caixa de teste ponta a ponta", True),
    ("fator_he", "Fator de hora extra na instalação e comissionamento", 1.5, "x",
     "Instalação fora do horário de separação (hora extra 50%). Use 1,0 se for feita no horário normal.", True),
    ("turnos", "Turnos de separação por dia", 2, "turnos", "CONFIRMAR com a operação (define pessoas a treinar)", True),
    ("h_treino_oper", "Treinamento por operador", 2, "h", "Estimativa: uso da estação, fechamento, falhas comuns", True),
    ("h_treino_manut", "Treinamento da equipe de manutenção (total)", 16, "h", "Estimativa: troca de módulo, diagnóstico, painel", True),
    ("reserva", "Reserva técnica (peças sobressalentes)", 0.10, "%", "Sobre módulos, nós e concentradores", True),
    ("frete", "Frete das compras nacionais", 0.05, "%",
     "Estimativa sobre as peças e a infraestrutura; itens importados já têm frete no fator de importação", True),
    ("contingencia", "Contingência do projeto", 0.10, "%",
     "Reserva para imprevistos de instalação e variação de preço (aplicada no Resumo)", True),
    ("m_barr", "Cabo de barramento por segmento (entre nós/módulos)", 3, "m", "Estimativa de layout", True),
    ("m_saida", "Cabo de cada saída do concentrador até o 1º equipamento", 6, "m",
     "Média das 2 saídas: alto giro ≈ 2 m (logo atrás) e baixo giro ≈ 10 m (por cima das esteiras)", True),
    ("larg_alto", "Largura de um prédio do alto giro", 0.25, "m", "Desenho de 26/10/2026: 2,5 m de alto giro por estação (10 prédios)", True),
    ("linhas_alto", "Linhas horizontais de canaleta na frente do alto giro", 2, "linhas",
     "1 acima do andar 3 (barramento dos nós) + 1 nos andares 4-5 (E+Q)", True),
    ("alt_vertical", "Descida vertical de canaleta por prédio do alto giro", 1.2, "m", "Do nó até os módulos Q dos andares 0-3", True),
    ("comp_corredor", "Comprimento de prateleira por lado do corredor", 2.0, "m",
     "Desenho de 26/10/2026: prateleiras de 50 cm × 2 m; 1 linha de canaleta na altura dos olhos em cada lado", True),
    ("m_travessia", "Travessia por estação (concentrador atrás → por cima das esteiras → corredor)", 6, "m",
     "Concentrador fica atrás da estação, do outro lado do alto giro", True),
    ("m_rede", "Cabo de rede do switch até a estação", 50, "m",
     "Switch instalado no meio de cada segmento (máx. 90 m por lance)", True),
    ("portas_switch", "Estações atendidas por switch de 24 portas", 22, "estações", "Deixa portas livres para uplink/reserva", True),
    ("m_uplink", "Cabo de uplink por switch até a rede corporativa", 80, "m",
     "Se passar de 90 m, trocar por fibra/conversor de mídia", True),
    ("pct_sem_tomada", "Estações sem tomada 220 V a até 3 m", 0.5, "%", "CONFIRMAR com a manutenção predial", True),
    ("frota", "Frota de caixas retornáveis (Arq. B)", 5000, "caixas",
     "Estimativa: ~1,7 × 3.000 volumes/dia. CONFIRMAR com a operação.", True),
    ("pulseiras_por_est", "Pulseiras por estação (Arq. B1)", 2.2, "un",
     "2 por estação (uma carregando; 500 mAh dura cerca de 5 h lendo) + 10% reserva", True),
    ("m_5v_extra", "Cabo de força extra por estação na variante 5 V", 30, "m", "Distribuição de 5 V exige mais pontos de injeção", True),
    ("m_cadeia_q", "Cabo por Módulo Q em cadeia (T1/T2)", 1.0, "m", "Módulos vizinhos ligados em sequência", True),
    ("h_final_ems", "Montagem final + gravação + teste por módulo vindo da EMS (T1)", 0.1, "h/un", "Estimativa", True),
]
REF = {k: "Premissas!$B$%d" % i for i, (k, *_r) in enumerate(premissas, 2)}
for i, (k, d, v, u, f, am) in enumerate(premissas, 2):
    formula = isinstance(v, str)
    if formula:
        v = v.format(**{kk: "$B$%d" % j for j, (kk, *_r) in enumerate(premissas, 2)})
    cel(wp, "A%d" % i, d)
    c = cel(wp, "B%d" % i, v, PRETO if formula else AZUL, fill=AMARELO if am else None)
    c.number_format = PCT if u == "%" else (NUM2 if isinstance(v, float) or formula else NUM)
    cel(wp, "C%d" % i, u)
    cel(wp, "D%d" % i, f, wrap=True)

# =====================================================================
# 3. ESTACOES
# =====================================================================
we = wb.create_sheet("Estações")
tit = ["Rua (estação)", "Segmento", "Tipo", "Piloto (S/N)", "Total de endereços",
       "Alto giro 0-3\n(Módulos Q)", "Alto giro 4-5\n(endereços)", "E+Q alto giro\n4-5 (2 zonas)",
       "Baixo giro\n(endereços)", "Zonas baixo giro\n(E+Q)", "Nós de coluna", "Prédios do\nbaixo giro",
       "Módulos E+Q", "Segmentos de\nbarramento", "Concentrador"]
cabecalho(we, 1, tit, [11, 14, 10, 10, 12, 13, 13, 13, 12, 13, 11, 12, 11, 13, 12])
ordem = sorted(EST)
for i, rua in enumerate(ordem, 2):
    e = EST[rua]
    piloto = "S" if rua in ("04", "27") else "N"      # estações da POC
    vals = [rua, e["segmento"], e["tipo"], piloto, e["total"], e["alto_0a3"], e["alto_45"],
            e["blocos_alto_45"], e["baixo"], e["zonas_baixo"], e["nos_coluna"], e["predios_fundo"]]
    for j, v in enumerate(vals, 1):
        c = we.cell(row=i, column=j, value=v)
        c.font = AZUL if j >= 4 else PRETO
        c.border = BORDA
        c.alignment = Alignment(horizontal="center")
        if j >= 5:
            c.number_format = NUM
    we.cell(row=i, column=13, value="=H%d+J%d" % (i, i))
    we.cell(row=i, column=14, value="=K%d+M%d" % (i, i))
    we.cell(row=i, column=15, value=1)
    for j in (13, 14, 15):
        c = we.cell(row=i, column=j)
        c.font = PRETO if j != 15 else AZUL
        c.border = BORDA
        c.alignment = Alignment(horizontal="center")
        c.number_format = NUM
ULT = len(ordem) + 1
LC, LP = ULT + 2, ULT + 3       # linhas de total completo / piloto
we.cell(row=LC, column=1, value="Total da operação completa").font = NEG
we.cell(row=LP, column=1, value="Total do piloto (Piloto = S)").font = NEG
dv = DataValidation(type="list", formula1='"S,N"', allow_blank=False)
we.add_data_validation(dv)
dv.add("D2:D%d" % ULT)
for j in range(5, 16):
    L = get_column_letter(j)
    c1 = we.cell(row=LC, column=j, value="=SUM(%s2:%s%d)" % (L, L, ULT))
    c2 = we.cell(row=LP, column=j, value='=SUMIF($D$2:$D$%d,"S",%s2:%s%d)' % (ULT, L, L, ULT))
    for c in (c1, c2):
        c.font = NEG
        c.fill = TOTAL
        c.number_format = NUM
        c.border = BORDA
        c.alignment = Alignment(horizontal="center")
we.cell(row=LC, column=4, value='=COUNTA(A2:A%d)' % ULT).number_format = NUM
we.cell(row=LP, column=4, value='=COUNTIF(D2:D%d,"S")' % ULT).number_format = NUM
for r in (LC, LP):
    we.cell(row=r, column=4).font = NEG
    we.cell(row=r, column=4).fill = TOTAL
we.cell(row=LP + 2, column=1,
        value="Coluna D (Piloto) = quantidade de estações nas linhas de total. "
              "Fonte: wms_endereco (cd_esteira 1 e 2), exportado em 23/10/2026; rua 90 (Marketing/Almoxarifado) desconsiderada.").font = NOTA
we.cell(row=LP + 3, column=1,
        value="Regras: Alto giro = prédios 001-010; andares 4-5 com 2 displays E+Q por estação (zonas 001-005 e 006-010). "
              "Baixo giro = demais prédios (corredor, 015 e esquinas): 2 zonas E+Q por lado do corredor; esquinas com zonas de 4 prédios.").font = NOTA
we.cell(row=LP + 4, column=1,
        value="Segmentos de barramento = 1 cabo de entrada por nó e por módulo E+Q (ligação em cadeia nas 2 portas do concentrador).").font = NOTA
we.freeze_panes = "B2"


def tot(col, cenario):
    """Referencia ao total da aba Estações: cenario 'C' completo, 'P' piloto."""
    linha = LC if cenario == "C" else LP
    return "'Estações'!$%s$%d" % (col, linha)


def par(tpl, extra=None):
    """Formula para piloto e completo: {F},{K}... = totais da aba Estações; {chave} = premissa; extra = {X: (piloto, completo)}."""
    out = []
    for k, cen in enumerate("PC"):
        d = {c: tot(c, cen) for c in "EFGHIJKLMNO"}
        d.update(REF)
        for nome, v in (extra or {}).items():
            d[nome] = v[k]
        out.append("=" + tpl.format(**d))
    return out


# =====================================================================
# 4. PRECOS
# =====================================================================
wpr = wb.create_sheet("Preços")
cabecalho(wpr, 1, ["Código", "Item", "Unidade", "Preço (R$)", "Preço (US$)", "Fonte / data", "Status"],
          [9, 56, 10, 13, 12, 72, 12])
C = "=E{r}*" + REF["cambio"] + "*" + REF["fator_imp"]
precos = [
    # cod, item, unidade, brl, usd, fonte, status
    ("P01", "Display 7 segmentos 4 dígitos TM1637 (usa 3 dígitos, quantidade)", "un", 9.98, None,
     "Usinainfo, Módulo Display TM1637 4 dígitos vermelho, 23/10/2026 (Pix)", "Cotado"),
    ("P02", "Display 7 segmentos 8 dígitos MAX7219 (endereço)", "un", 20.17, None,
     "Usinainfo, Display MAX7219 8 dígitos SPI, 23/10/2026 (Pix)", "Cotado"),
    ("P03", "Microcontrolador Raspberry Pi Pico (RP2040)", "un", 39.90, None,
     "MakerHero, Raspberry Pi Pico, 23/10/2026", "Cotado"),
    ("P04", "Módulo transceptor RS-485 (MAX485)", "un", 8.90, None,
     "Saravati, Módulo conversor TTL/RS485 MAX485, 23/10/2026", "Cotado"),
    ("P05", "Placa do concentrador RP2040 + Ethernet W5500 (W5500-EVB-Pico)", "un", None, 12.50,
     "SparkFun DEV-19958 US$ 12,50, 23/10/2026; convertido com câmbio × fator de importação", "Cotado (US$)"),
    ("P06", "Fonte Mean Well LRS-150-12 (12 V / 12,5 A / 150 W)", "un", 189.98, None,
     "MCEIG (fontesmeanwell.com.br), 23/10/2026. Carga estimada < 40 W por estação: folga para expansão", "Cotado"),
    ("P07", "Fonte Mean Well LRS-150-5 (5 V / 22 A), variante 5 V", "un", 185.00, None,
     "Estimativa (mesma família da LRS-150-12), cotar", "Estimativa"),
    ("P08", "Regulador step-down 12 V → 5 V (mini, 3 A)", "un", 8.00, None, "Estimativa de mercado, cotar", "Estimativa"),
    ("P09", "Botão de confirmação metálico 16 mm com anel de LED 12 V (momentâneo)", "un", 18.00, None,
     "Estimativa de mercado, cotar. Botão plástico comum não aguenta ~6.000 toques/dia por estação", "Estimativa"),
    ("P10", "Conector RJ45 fêmea para placa (jack)", "un", 2.50, None, "Estimativa, cotar", "Estimativa"),
    ("P11", "Plug RJ45 Cat5e (macho)", "un", 0.80, None, "Estimativa (pote com 100), cotar", "Estimativa"),
    ("P12", "Cabo Cat5e (por metro)", "m", "=1229.41/305", None,
     "KaBuM, Cabo Cat5e Fortrek FK500C 305 m R$ 1.229,41 (Pix), 23/10/2026", "Cotado"),
    ("P13", "Carcaça do Módulo Q (impressão 3D PETG ou caixa plástica pequena)", "un", 5.00, None, "Estimativa", "Estimativa"),
    ("P14", "Carcaça do Módulo E+Q / Nó de coluna", "un", 15.00, None, "Estimativa", "Estimativa"),
    ("P15", "Placa base universal + bornes (nó / módulo E+Q)", "un", 10.00, None, "Estimativa", "Estimativa"),
    ("P16", "Gabinete do concentrador ventilado (caixa IP54 + cooler 40 mm 12 V + filtro)", "un", 90.00, None,
     "Estimativa, cotar (ambiente com poeira e calor)", "Estimativa"),
    ("P17", "Trilho DIN + bornes + disjuntor 2 A", "kit", 45.00, None, "Estimativa, cotar", "Estimativa"),
    ("P18", "Canaleta / perfil de fixação dos módulos (por metro)", "m", 9.00, None,
     "Estimativa (canaleta PVC 20×10 mm em barra de 2 m), cotar", "Estimativa"),
    ("P19", "Fixações diversas por estação (abraçadeiras, parafusos, etiquetas, plugs da rede)", "kit", 40.00, None,
     "Estimativa", "Estimativa"),
    ("P20", "Switch 24 portas Gigabit TP-Link TL-SG1024DE", "un", 769.32, None, "KaBuM, 23/10/2026 (Pix)", "Cotado"),
    ("P21", "(não usado) LED guia de coluna, retirado em 26/10/2026", "LED", 0, None, "Não faz parte do projeto", "Retirado"),
    ("P22", "Tag NFC NTAG213 adesiva 13,56 MHz", "un", 2.37, None,
     "Loja RFID, lote 100 un R$ 237,41 à vista, 23/10/2026 (lote de 5.000 deve sair mais barato)", "Cotado"),
    ("P23", "Suporte/proteção da tag na caixa retornável", "un", 0.50, None, "Estimativa", "Estimativa"),
    ("P24", "ESP32-C3 Super Mini (pulseira)", "un", 29.89, None, "Usinainfo, 23/10/2026 (Pix)", "Cotado"),
    ("P25", "Módulo leitor NFC PN532 (43×40 mm, com cartão e chaveiro)", "un", 42.65, None,
     "MakerHero, Kit Módulo Leitor RFID NFC PN532, 23/10/2026 (Pix)", "Cotado"),
    ("P26", "Kit energia pulseira (LiPo 500 mAh + TP4056 + motor vibração + botão)", "kit", 35.00, None, "Estimativa, cotar", "Estimativa"),
    ("P27", "Carcaça 3D + pulseira velcro", "un", 20.00, None, "Estimativa", "Estimativa"),
    ("P28", "Display LCD 20×4 I2C (painel da estação, Arq. B2)", "un", 35.00, None, "Estimativa, cotar", "Estimativa"),
    ("P29", "Módulo PTL comercial 3 dígitos + botão (RS-485)", "un", None, 30.00,
     "Alibaba, Sophia T-P3: US$ 32 (1-299 un) / US$ 28 (300+), 23/10/2026; média US$ 30", "Cotado (US$)"),
    ("P30", "Módulo PTL comercial endereço + quantidade (display alfanumérico)", "un", None, 65.00,
     "Estimativa (não há preço público), solicitar cotação a fornecedores", "Estimativa"),
    ("P31", "Controlador/gateway PTL comercial por estação", "un", None, 220.00,
     "Estimativa (não há preço público), solicitar cotação", "Estimativa"),
    ("P32", "PCB própria montada por EMS, Módulo Q (MCU + display 0,56\" + botão + 2× RJ45)", "un", None, 9.00,
     "Estimativa com base na tabela de montagem JLCPCB (jlcpcb.com/help/article/pcb-assembly-price) + componentes", "Estimativa"),
    ("P33", "PCB própria montada por EMS, Módulo E+Q", "un", None, 22.00, "Estimativa (idem)", "Estimativa"),
    ("P34", "PCB própria montada por EMS, Concentrador (MCU + Ethernet + 2× RS-485)", "un", None, 35.00, "Estimativa (idem)", "Estimativa"),
    ("P35", "Engenharia de projeto das PCBs (terceirizada, valor único)", "projeto", 15000.00, None,
     "Estimativa, solicitar orçamento a design house", "Estimativa"),
    ("P36", "Leitor RFID UHF de pulso comercial (FEIG HyWEAR Compact), referência", "un", None, 997.00,
     "atlasRFIDstore, US$ 997, 22/10/2026 (UHF; não lê tags NFC)", "Cotado (US$)"),
    ("P37", "Leitor HF/NFC fixo industrial (USB/RS-232), alternativa comercial ao PN532", "un", 350.00, None,
     "Estimativa, cotar (ex.: ACS ACR122U ou similar)", "Estimativa"),
    ("P38", "Cabo PP 3×1,0 mm² (por metro)", "m", "=913.31/100", None,
     "KaBuM, Cabo PP 3×1,00 mm² flexível preto 100 m R$ 913,31 (Pix), 26/10/2026", "Cotado"),
    ("P39", "Plugue 2P+T 10 A + prensa-cabo", "un", 12.00, None, "Estimativa, cotar", "Estimativa"),
    ("P40", "DPS / protetor de surto 275 V para o concentrador", "un", 35.00, None, "Estimativa, cotar", "Estimativa"),
    ("P41", "Kit rede: alicate de crimpar + testador de cabo + 40 m Cat5e Nexans + 10 RJ45", "kit", 285.99, None,
     "KaBuM, produto 386101, 26/10/2026 (Pix)", "Cotado"),
    ("P42", "Bancada de eletrônica: estação de solda + multímetro + consumíveis", "kit", 700.00, None,
     "Estimativa, cotar", "Estimativa"),
    ("P43", "Nobreak 720 VA bivolt (Intelbras XNB) para o switch", "un", 368.10, None,
     "KaBuM via escorregaopreco.com.br, oferta atualizada ~28 dias antes de 26/10/2026", "Cotado"),
    ("P44", "Mini rack de parede 5U 19\" + patch cords + organizador", "kit", 550.00, None,
     "Estimativa (ex.: Intelbras MRM 537), cotar", "Estimativa"),
    ("P45", "Carregador USB 2 portas para pulseiras (Arq. B1)", "un", 35.00, None, "Estimativa, cotar", "Estimativa"),
    ("P46", "Ponto de tomada 220 V novo (material + eletricista)", "un", 350.00, None,
     "Estimativa, cotar com a manutenção predial", "Estimativa"),
    ("P48", "Painel de MDF 1,2 × 0,6 m com longarina para a bancada + parafusos", "kit", 150.00, None, "Estimativa, cotar", "Estimativa"),
    ("P49", "Switch 5 portas (bancada)", "un", 90.00, None, "Estimativa, cotar (ou porta livre de switch existente)", "Estimativa"),
    ("P47", "Sinal luminoso de baixo giro (sinaleiro LED âmbar 12 V 22 mm + haste + transistor no nó)", "un", 45.00, None,
     "Estimativa, cotar", "Estimativa"),
]
PRECO = {}
for i, (cod, item, un, brl, usd, fonte, st) in enumerate(precos, 2):
    PRECO[cod] = "'Preços'!$D$%d" % i
    cel(wpr, "A%d" % i, cod)
    cel(wpr, "B%d" % i, item, wrap=True)
    cel(wpr, "C%d" % i, un)
    if brl is None:
        cel(wpr, "D%d" % i, C.format(r=i), PRETO, BRL)
    else:
        cel(wpr, "D%d" % i, brl, AZUL if not str(brl).startswith("=") else PRETO, BRL,
            fill=AMARELO if st == "Estimativa" else None)
    if usd is not None:
        cel(wpr, "E%d" % i, usd, AZUL, '"US$ "#,##0.00', fill=AMARELO if st == "Estimativa" else None)
    else:
        cel(wpr, "E%d" % i, None)
    cel(wpr, "F%d" % i, fonte, wrap=True)
    cel(wpr, "G%d" % i, st, fill=AMARELO if st == "Estimativa" else None)
wpr.freeze_panes = "C2"

# =====================================================================
# 5. COMPOSICAO (BOM)
# =====================================================================
wc = wb.create_sheet("Composição")
cabecalho(wc, 1, ["Conjunto", "Código", "Item", "Qtd por unidade", "Preço unit. (R$)", "Subtotal (R$)"],
          [26, 9, 60, 14, 15, 15])
alimentacao = [("P38", 3), ("P39", 1), ("P40", 1)]       # cabo até a tomada, plugue, DPS
bom = {
    "Módulo Q": [("P01", 1), ("P09", 1), ("P13", 1), ("P10", 1), ("P12", 2), ("P11", 2)],
    "Nó de coluna": [("P03", 1), ("P04", 1), ("P08", 1), ("P15", 1), ("P10", 6), ("P14", 1)],
    "Módulo E+Q": [("P02", 1), ("P01", 1), ("P03", 1), ("P04", 1), ("P08", 1), ("P09", 1), ("P10", 2),
                   ("P14", 1), ("P15", 1)],
    "Concentrador 12 V": [("P05", 1), ("P04", 2), ("P06", 1), ("P08", 1), ("P16", 1), ("P17", 1), ("P10", 2)] + alimentacao,
    "Concentrador 5 V": [("P05", 1), ("P04", 2), ("P07", 1), ("P16", 1), ("P17", 1), ("P10", 2)] + alimentacao,
    "Pulseira RFID (B1)": [("P24", 1), ("P25", 1), ("P26", 1), ("P27", 1)],
    "Leitor fixo + painel (B2)": [("P25", 1), ("P28", 1), ("P09", 1), ("P03", 1), ("P04", 1), ("P08", 1), ("P15", 1),
                                  ("P10", 2), ("P14", 1), ("P12", 5), ("P11", 2)],
    "Tag por caixa (B)": [("P22", 1), ("P23", 1)],
}
r = 2
UNIT = {}
for conj, itens in bom.items():
    ini = r
    for cod, q in itens:
        cel(wc, "A%d" % r, conj)
        cel(wc, "B%d" % r, cod)
        cel(wc, "C%d" % r, "=INDEX('Preços'!$B:$B,MATCH(B%d,'Preços'!$A:$A,0))" % r, VERDE, wrap=True)
        cel(wc, "D%d" % r, q, AZUL, NUM2)
        cel(wc, "E%d" % r, "=INDEX('Preços'!$D:$D,MATCH(B%d,'Preços'!$A:$A,0))" % r, VERDE, BRL)
        cel(wc, "F%d" % r, "=D%d*E%d" % (r, r), PRETO, BRL)
        r += 1
    faixa_total(wc, r, "Custo unitário, " + conj)
    cel(wc, "F%d" % r, "=SUM(F%d:F%d)" % (ini, r - 1), NEG, BRL, fill=TOTAL)
    UNIT[conj] = "'Composição'!$F$%d" % r
    r += 2
for t in ("Variante 5 V: os módulos e nós não usam regulador step-down (P08); o custo é descontado na aba A-Interno.",
          "Concentrador: inclui alimentação até a tomada (3 m de cabo PP, plugue) e protetor de surto.",
          "Leitor fixo (B2): tem microcontrolador e RS-485 próprios e entra no barramento como mais um nó (cabo de 5 m)."):
    cel(wc, "A%d" % r, t, NOTA, borda=False)
    r += 1

# =====================================================================
# 6. A-INTERNO
# =====================================================================
wa = wb.create_sheet("A-Interno")
cabecalho(wa, 1, ["Item", "Base de cálculo", "Qtd piloto", "Qtd completo", "Preço unit. (R$)",
                  "Total piloto (R$)", "Total completo (R$)"], [46, 56, 12, 13, 15, 17, 18])
cel(wa, "A2", "Arquitetura A, produção e montagem interna (12 V)", SUB, borda=False)
res = "(1+{reserva})"
LIN = {}       # nome curto -> linha


def linha_item(ws, r, nome, base, qp, qc, preco, fmt=NUM):
    cel(ws, "A%d" % r, nome, wrap=True)
    cel(ws, "B%d" % r, base, wrap=True)
    cel(ws, "C%d" % r, qp, PRETO if str(qp).startswith("=") else AZUL, fmt)
    cel(ws, "D%d" % r, qc, PRETO if str(qc).startswith("=") else AZUL, fmt)
    cel(ws, "E%d" % r, "=" + preco, VERDE, BRL)
    cel(ws, "F%d" % r, "=C%d*E%d" % (r, r), PRETO, BRL)
    cel(ws, "G%d" % r, "=D%d*E%d" % (r, r), PRETO, BRL)


modulos = [
    ("q", "Módulos Q (alto giro 0-3)", "endereços alto giro 0-3 × (1 + reserva)", "ROUNDUP({F}*%s,0)" % res, UNIT["Módulo Q"]),
    ("no", "Nós de coluna (alto giro)", "prédios do alto giro × (1 + reserva)", "ROUNDUP({K}*%s,0)" % res, UNIT["Nó de coluna"]),
    ("eq", "Módulos E+Q (A4/A5 + zonas baixo giro)", "displays E+Q × (1 + reserva)", "ROUNDUP({M}*%s,0)" % res, UNIT["Módulo E+Q"]),
    ("conc", "Concentradores 12 V", "1 por estação × (1 + reserva)", "ROUNDUP({O}*%s,0)" % res, UNIT["Concentrador 12 V"]),
    ("barr", "Cabo de barramento Cat5e (m)", "segmentos × m por segmento + 2 saídas × m de saída por estação",
     "{N}*{m_barr}+{O}*2*{m_saida}", PRECO["P12"]),
    ("plug", "Plugs RJ45 do barramento", "2 por segmento + 4 por estação (saídas)", "{N}*2+{O}*4", PRECO["P11"]),
]
infra = [
    ("can", "Canaleta / perfil de fixação (m)",
     "prédios alto giro × (largura × linhas + descida) + 2 lados × corredor + travessia, por estação",
     "{K}*({larg_alto}*{linhas_alto}+{alt_vertical})+{O}*2*{comp_corredor}+{O}*{m_travessia}", PRECO["P18"]),
    ("fix", "Fixações diversas", "1 kit por estação", "{O}", PRECO["P19"]),
    ("sin", "Sinal luminoso de baixo giro", "1 por estação, ao lado da posição de pega", "{O}", PRECO["P47"]),
    ("rede", "Cabo de rede até a estação (m)", "metros por estação", "{O}*{m_rede}", PRECO["P12"]),
    ("sw", "Switch 24 portas", "estações ÷ portas por switch (arredondado para cima)", "ROUNDUP({O}/{portas_switch},0)", PRECO["P20"]),
    ("nb", "Nobreak do switch", "1 por switch", "{SW}", PRECO["P43"]),
    ("rack", "Mini rack de parede do switch", "1 por switch", "{SW}", PRECO["P44"]),
    ("upl", "Cabo de uplink até a rede corporativa (m)", "switches × m de uplink", "{SW}*{m_uplink}", PRECO["P12"]),
    ("tom", "Pontos de tomada 220 V novos", "estações × % sem tomada próxima (arredondado para cima)",
     "ROUNDUP({O}*{pct_sem_tomada},0)", PRECO["P46"]),
    ("ferr1", "Ferramentas de rede (crimpar e testar)", "1 kit (compra única)", "1", PRECO["P41"]),
    ("ferr2", "Bancada de eletrônica (solda, multímetro)", "1 kit (compra única)", "1", PRECO["P42"]),
]
r = 3
for chave, nome, base, tpl, preco in modulos + infra:
    LIN[chave] = r
    extra = {"SW": ("C%d" % LIN["sw"], "D%d" % LIN["sw"])} if "{SW}" in tpl else None
    qp, qc = par(tpl, extra)
    linha_item(wa, r, nome, base, qp, qc, preco)
    r += 1
INF_INI, INF_FIM = LIN["can"], r - 1
LIN["frete"] = r
cel(wa, "A%d" % r, "Frete das compras")
cel(wa, "B%d" % r, "% sobre peças e infraestrutura acima")
cel(wa, "C%d" % r, None)
cel(wa, "D%d" % r, None)
cel(wa, "E%d" % r, "=" + REF["frete"], VERDE, PCT)
cel(wa, "F%d" % r, "=E%d*SUM(F3:F%d)" % (r, r - 1), PRETO, BRL)
cel(wa, "G%d" % r, "=E%d*SUM(G3:G%d)" % (r, r - 1), PRETO, BRL)
r += 1
A_PECAS = r
faixa_total(wa, r, "SUBTOTAL MATERIAIS, Arquitetura A (12 V)")
cel(wa, "F%d" % r, "=SUM(F3:F%d)" % (r - 1), NEG, BRL, fill=TOTAL)
cel(wa, "G%d" % r, "=SUM(G3:G%d)" % (r - 1), NEG, BRL, fill=TOTAL)
r += 2
cel(wa, "A%d" % r, "Mão de obra interna e treinamento", SUB, borda=False)
r += 1
can = ("C%d" % LIN["can"], "D%d" % LIN["can"])


def qtd(chave):
    return ("C%d" % LIN[chave], "D%d" % LIN[chave])


mo = [
    ("Montagem Módulos Q (h)", "módulos × h/un", "{X}*{h_mod_q}", {"X": qtd("q")}, REF["custo_hora"]),
    ("Montagem Nós de coluna (h)", "nós × h/un", "{X}*{h_no}", {"X": qtd("no")}, REF["custo_hora"]),
    ("Montagem Módulos E+Q (h)", "módulos × h/un", "{X}*{h_mod_eq}", {"X": qtd("eq")}, REF["custo_hora"]),
    ("Montagem Concentradores (h)", "concentradores × h/un", "{X}*{h_conc}", {"X": qtd("conc")}, REF["custo_hora"]),
    ("Instalação nas estações (h)",
     "(fixo por estação + Q + nós + E+Q + canaleta) × fator hora extra",
     "({O}*{h_inst_est}+{F}*{h_inst_q}+{K}*{h_inst_no}+{M}*{h_inst_eq}+{CAN}*{h_inst_m})*{fator_he}", {"CAN": can},
     REF["custo_hora"]),
    ("Testes e comissionamento (h)", "estações × h/estação × fator hora extra", "{O}*{h_comiss}*{fator_he}", None,
     REF["custo_hora"]),
    ("Treinamento dos operadores (h)", "estações × turnos × h por operador", "{O}*{turnos}*{h_treino_oper}", None,
     REF["custo_hora_oper"]),
    ("Treinamento da manutenção (h)", "total da equipe", "{h_treino_manut}", None, REF["custo_hora"]),
]
ini = r
for nome, base, tpl, extra, preco in mo:
    qp, qc = par(tpl, extra)
    linha_item(wa, r, nome, base, qp, qc, preco, NUM2)
    if nome.startswith("Treinamento dos"):
        LIN["treino_ini"] = r
    r += 1
LIN["treino_fim"] = r - 1
A_MO = r
faixa_total(wa, r, "SUBTOTAL MÃO DE OBRA E TREINAMENTO", "BE")
cel(wa, "C%d" % r, "=SUM(C%d:C%d)" % (ini, r - 1), NEG, NUM2, fill=TOTAL)
cel(wa, "D%d" % r, "=SUM(D%d:D%d)" % (ini, r - 1), NEG, NUM2, fill=TOTAL)
cel(wa, "F%d" % r, "=SUM(F%d:F%d)" % (ini, r - 1), NEG, BRL, fill=TOTAL)
cel(wa, "G%d" % r, "=SUM(G%d:G%d)" % (ini, r - 1), NEG, BRL, fill=TOTAL)
r += 2
A_TOT = r
faixa_total(wa, r, "TOTAL, Arquitetura A interna (materiais + mão de obra)")
cel(wa, "F%d" % r, "=F%d+F%d" % (A_PECAS, A_MO), NEG, BRL, fill=TOTAL)
cel(wa, "G%d" % r, "=G%d+G%d" % (A_PECAS, A_MO), NEG, BRL, fill=TOTAL)
r += 2
cel(wa, "A%d" % r, "Variante 5 V (materiais)", SUB, borda=False)
r += 1
v5 = [
    ("(−) Reguladores step-down dispensados", "nós + módulos E+Q + concentradores",
     "-({A}+{B}+{Cc})", {"A": qtd("no"), "B": qtd("eq"), "Cc": qtd("conc")}, PRECO["P08"]),
    ("(±) Fonte 5 V no lugar da 12 V", "concentradores", "{A}", {"A": qtd("conc")}, "(%s-%s)" % (PRECO["P07"], PRECO["P06"])),
    ("(+) Cabo de força extra para injeção de 5 V (m)", "metros por estação", "{O}*{m_5v_extra}", None, PRECO["P38"]),
]
ini = r
for nome, base, tpl, extra, preco in v5:
    qp, qc = par(tpl, extra)
    linha_item(wa, r, nome, base, qp, qc, preco)
    r += 1
A_5V = r
faixa_total(wa, r, "SUBTOTAL MATERIAIS, variante 5 V")
cel(wa, "F%d" % r, "=F%d+SUM(F%d:F%d)" % (A_PECAS, ini, r - 1), NEG, BRL, fill=TOTAL)
cel(wa, "G%d" % r, "=G%d+SUM(G%d:G%d)" % (A_PECAS, ini, r - 1), NEG, BRL, fill=TOTAL)
r += 2

INFRA_F = "SUM('A-Interno'!F%d:F%d)" % (INF_INI, INF_FIM)
INFRA_G = "SUM('A-Interno'!G%d:G%d)" % (INF_INI, INF_FIM)
TREINO_F = "SUM('A-Interno'!F%d:F%d)" % (LIN["treino_ini"], LIN["treino_fim"])
TREINO_G = "SUM('A-Interno'!G%d:G%d)" % (LIN["treino_ini"], LIN["treino_fim"])
CAN_REF = ("'A-Interno'!C%d" % LIN["can"], "'A-Interno'!D%d" % LIN["can"])

# =====================================================================
# 7. A-TERCEIROS
# =====================================================================
wt = wb.create_sheet("A-Terceiros")
cabecalho(wt, 1, ["Item", "Base de cálculo", "Qtd piloto", "Qtd completo", "Preço unit. (R$)",
                  "Total piloto (R$)", "Total completo (R$)"], [46, 50, 12, 13, 15, 17, 18])
alim = "+".join("%s*%s" % (PRECO[c], q) for c, q in alimentacao)


def linha_ref(ws, r, nome, base, f, g):
    cel(ws, "A%d" % r, nome, wrap=True)
    cel(ws, "B%d" % r, base, wrap=True)
    for col in "CDE":
        cel(ws, "%s%d" % (col, r), None)
    cel(ws, "F%d" % r, "=" + f, VERDE, BRL)
    cel(ws, "G%d" % r, "=" + g, VERDE, BRL)


inst_sem_no = "({O}*{h_inst_est}+{F}*{h_inst_q}+{M}*{h_inst_eq}+{CAN}*{h_inst_m}+{O}*{h_comiss})*{fator_he}"
cabo_cadeia = "{F}*{m_cadeia_q}+{M}*{m_barr}+{O}*2*{m_saida}"


def bloco_terceiros(ws, r, titulo, itens, rot_total):
    cel(ws, "A%d" % r, titulo, SUB, borda=False)
    r += 1
    ini = r
    for it in itens:
        if it[0] == "ref":
            linha_ref(ws, r, it[1], it[2], it[3], it[4])
        else:
            nome, base, tpl, extra, preco = it
            if callable(extra):
                extra = extra(ini)
            qp, qc = par(tpl, extra)
            linha_item(ws, r, nome, base, qp, qc, preco, NUM2 if "(h)" in nome else NUM)
        r += 1
    faixa_total(ws, r, rot_total)
    cel(ws, "F%d" % r, "=SUM(F%d:F%d)" % (ini, r - 1), NEG, BRL, fill=TOTAL)
    cel(ws, "G%d" % r, "=SUM(G%d:G%d)" % (ini, r - 1), NEG, BRL, fill=TOTAL)
    return r


t1 = [
    ("Módulo Q, PCB montada + carcaça", "endereços alto giro 0-3 × (1 + reserva)", "ROUNDUP({F}*%s,0)" % res, None,
     "%s+%s" % (PRECO["P32"], PRECO["P13"])),
    ("Módulo E+Q, PCB montada + carcaça + botão", "displays E+Q × (1 + reserva)", "ROUNDUP({M}*%s,0)" % res, None,
     "%s+%s+%s" % (PRECO["P33"], PRECO["P14"], PRECO["P09"])),
    ("Concentrador, PCB montada + fonte + gabinete + alimentação", "estações × (1 + reserva)", "ROUNDUP({O}*%s,0)" % res, None,
     "%s+%s+%s+%s+%s" % (PRECO["P34"], PRECO["P06"], PRECO["P16"], PRECO["P17"], alim)),
    ("Cabo do barramento em cadeia (m)", "Q × m em cadeia + E+Q × m por segmento + saídas", cabo_cadeia, None, PRECO["P12"]),
    ("Plugs RJ45 do barramento", "2 por módulo + 4 por estação", "({F}+{M})*2+{O}*4", None, PRECO["P11"]),
    ("ref", "Infraestrutura (canaleta, rede, switches, energia, ferramentas)", "igual à Arquitetura A", INFRA_F, INFRA_G),
    ("Engenharia de projeto das PCBs (valor único)", "1 projeto", "1", None, PRECO["P35"]),
    ("Montagem final + gravação + teste (h)", "módulos e concentradores recebidos × h/un",
     "({A}+{B}+{Cc})*{h_final_ems}", lambda i: {"A": ("C%d" % i, "D%d" % i), "B": ("C%d" % (i + 1), "D%d" % (i + 1)),
                                                 "Cc": ("C%d" % (i + 2), "D%d" % (i + 2))}, REF["custo_hora"]),
    ("Instalação + comissionamento (h)", "como na Arquitetura A, sem nós de coluna", inst_sem_no, {"CAN": CAN_REF},
     REF["custo_hora"]),
    ("ref", "Treinamento (operadores e manutenção)", "igual à Arquitetura A", TREINO_F, TREINO_G),
]
T1_TOT = bloco_terceiros(wt, 2, "T1, Placas próprias fabricadas e montadas por EMS (cada módulo com seu microcontrolador)",
                         t1, "TOTAL T1, PCB própria (EMS)")
t2 = [
    ("Módulos PTL 3 dígitos (alto giro 0-3)", "endereços × (1 + reserva)", "ROUNDUP({F}*%s,0)" % res, None, PRECO["P29"]),
    ("Módulos PTL endereço + quantidade", "displays E+Q × (1 + reserva)", "ROUNDUP({M}*%s,0)" % res, None, PRECO["P30"]),
    ("Controlador/gateway PTL por estação", "estações", "{O}", None, PRECO["P31"]),
    ("Fonte 12 V + gabinete + alimentação por estação", "estações", "{O}", None,
     "%s+%s+%s+%s" % (PRECO["P06"], PRECO["P16"], PRECO["P17"], alim)),
    ("Cabo do barramento em cadeia (m)", "Q × m em cadeia + E+Q × m por segmento + saídas", cabo_cadeia, None, PRECO["P12"]),
    ("Plugs RJ45 do barramento", "2 por módulo + 4 por estação", "({F}+{M})*2+{O}*4", None, PRECO["P11"]),
    ("ref", "Infraestrutura (canaleta, rede, switches, energia, ferramentas)", "igual à Arquitetura A", INFRA_F, INFRA_G),
    ("Instalação + comissionamento (h)", "como na Arquitetura A, sem nós de coluna", inst_sem_no, {"CAN": CAN_REF},
     REF["custo_hora"]),
    ("ref", "Treinamento (operadores e manutenção)", "igual à Arquitetura A", TREINO_F, TREINO_G),
]
T2_TOT = bloco_terceiros(wt, T1_TOT + 2, "T2, Sistema PTL comercial importado (hardware pronto; software de integração continua interno)",
                         t2, "TOTAL T2, PTL comercial importado")
cel(wt, "A%d" % (T2_TOT + 2), "Valores em US$ convertidos por câmbio × fator de importação (aba Premissas). "
    "Itens de PTL comercial sem preço público devem ser cotados com fornecedores (ex.: ATOP, Aioi, fornecedores nacionais). "
    "Frete nacional não aplicado aqui: os itens principais são importados (frete no fator de importação).", NOTA, borda=False)

# =====================================================================
# 8. B-RFID
# =====================================================================
wb_ = wb.create_sheet("B-RFID")
cabecalho(wb_, 1, ["Item", "Base de cálculo", "Qtd piloto", "Qtd completo", "Preço unit. (R$)",
                   "Total piloto (R$)", "Total completo (R$)"], [48, 48, 12, 13, 15, 17, 18])
cel(wb_, "A2", "Arquitetura B = infraestrutura de displays da Arquitetura A + identificação da caixa por RFID", SUB, borda=False)
linha_ref(wb_, 3, "Infraestrutura de displays (Arquitetura A interna, materiais + mão de obra)", "aba A-Interno",
          "'A-Interno'!F%d" % A_TOT, "'A-Interno'!G%d" % A_TOT)
linha_item(wb_, 4, "Tags NFC nas caixas retornáveis (frota)", "frota de caixas (piloto: proporcional às estações)",
           "=ROUNDUP(%s*%s/%s,0)" % (REF["frota"], tot("O", "P"), tot("O", "C")), "=%s" % REF["frota"], UNIT["Tag por caixa (B)"])
linha_item(wb_, 5, "Associação tag ↔ caixa (cadastro inicial) (h)", "30 s por caixa", "=C4*30/3600", "=D4*30/3600",
           REF["custo_hora"], NUM2)
cel(wb_, "A7", "B1, Pulseira NFC (montagem própria)", SUB, borda=False)
qp, qc = par("ROUNDUP({O}*{pulseiras_por_est},0)")
linha_item(wb_, 8, "Pulseiras NFC (ESP32-C3 + PN532 + bateria)", "estações × pulseiras por estação", qp, qc, UNIT["Pulseira RFID (B1)"])
qp, qc = par("{O}")
linha_item(wb_, 9, "Carregador USB das pulseiras", "1 por estação", qp, qc, PRECO["P45"])
linha_item(wb_, 10, "Montagem das pulseiras (h)", "1,5 h por pulseira", "=C8*1.5", "=D8*1.5", REF["custo_hora"], NUM2)
faixa_total(wb_, 11, "TOTAL B1, displays + tags + pulseiras")
cel(wb_, "F11", "=SUM(F3:F5)+SUM(F8:F10)", NEG, BRL, fill=TOTAL)
cel(wb_, "G11", "=SUM(G3:G5)+SUM(G8:G10)", NEG, BRL, fill=TOTAL)
B1_TOT = 11
cel(wb_, "A13", "B2, Leitor fixo na estação + painel (sem item no corpo do operador)", SUB, borda=False)
qp, qc = par("{O}")
linha_item(wb_, 14, "Leitor PN532 + LCD 20×4 + botão Fechar (nó do barramento)", "1 por estação", qp, qc,
           UNIT["Leitor fixo + painel (B2)"])
linha_item(wb_, 15, "Montagem e instalação do leitor (h)", "2 h por estação", "=C14*2", "=D14*2", REF["custo_hora"], NUM2)
faixa_total(wb_, 16, "TOTAL B2, displays + tags + leitor fixo")
cel(wb_, "F16", "=SUM(F3:F5)+SUM(F14:F15)", NEG, BRL, fill=TOTAL)
cel(wb_, "G16", "=SUM(G3:G5)+SUM(G14:G15)", NEG, BRL, fill=TOTAL)
B2_TOT = 16
linha_item(wb_, 18, "Referência de mercado, leitor de pulso comercial", "FEIG HyWEAR (UHF) por pulseira, não lê NFC; exigiria tags UHF",
           "=C8", "=D8", PRECO["P36"])
cel(wb_, "A20", "Coletores MovFast já existentes não entram no custo (Arquitetura A). Na Arquitetura B eles deixam de ser "
    "necessários na separação, mas continuam disponíveis para as demais operações.", NOTA, borda=False)

# =====================================================================
# 9. SOFTWARE
# =====================================================================
wsw = wb.create_sheet("Software")
cabecalho(wsw, 1, ["Componente", "Arquitetura", "Horas estimadas", "Custo (R$)", "Observação"], [48, 14, 15, 16, 72])
sw = [
    ("Serviço PTL (API REST + MQTT + regras de separação)", "A e B", 240, "Núcleo: abre/fecha caixa, divide por estação/zona, ordena pela rota"),
    ("Integração Sybase ASE (tabelas PTL, views, procedures)", "A e B", 120, "Leitura de itens da caixa e gravação de eventos/log"),
    ("Integração API do sistema BHT (informações da caixa)", "A e B", 60, "Consulta da caixa e, se possível, liberação para a próxima estação"),
    ("App coletor Android (A1, app próprio)", "A", 160, "Leitura com gatilho para abrir/fechar, progresso da caixa"),
    ("Tela no WMS PowerBuilder para o coletor (A2)", "A2", 80, "Alternativa ao app; grava a leitura em tabela consumida pelo serviço"),
    ("Firmware do concentrador (Ethernet + MQTT + 2 barramentos RS-485)", "A e B", 120, "Base no protótipo MicroPython já validado"),
    ("Firmware dos nós de coluna e módulos E+Q", "A e B", 100, "Protocolo de barramento, endereçamento, auto-teste"),
    ("Leitura RFID no concentrador + painel da estação (B2) / pulseira (B1)", "B", 80, "Somente Arquitetura B"),
    ("Painel de monitoramento (gestão e TI)", "A e B", 120, "Status das estações, módulos com falha, produtividade"),
    ("Interface de controle: qual LED falhou + cadastro de LED por endereço (tela + botão)", "A e B", 40,
     "Diagnóstico no firmware (sem resposta, display sem ACK, botão travado), modo cadastro e tabela ptl_mapa_led"),
    ("Implantação do servidor (VM, broker MQTT, backup, monitoração)", "A e B", 24,
     "Considera VM em infraestrutura já existente; sem compra de servidor"),
    ("Testes integrados, piloto e homologação", "A e B", 120, "Inclui ajustes na POC (estações 04 e 27)"),
    ("Documentação e treinamento", "A e B", 40, "Operação, manutenção e suporte"),
]
for i, (comp, arq, h, obs) in enumerate(sw, 2):
    cel(wsw, "A%d" % i, comp, wrap=True)
    cel(wsw, "B%d" % i, arq)
    cel(wsw, "C%d" % i, h, AZUL, NUM, fill=AMARELO)
    cel(wsw, "D%d" % i, "=C%d*%s" % (i, REF["custo_hora_dev"]), PRETO, BRL0)
    cel(wsw, "E%d" % i, obs, wrap=True)
n = len(sw) + 2
rng = "B2:B%d" % (n - 1), "C2:C%d" % (n - 1)
faixa_total(wsw, n, "TOTAL, Arquitetura A com app próprio (A1)", "B")
cel(wsw, "C%d" % n, '=SUMIF(%s,"A e B",%s)+SUMIF(%s,"A",%s)' % (rng + rng), NEG, NUM, fill=TOTAL)
cel(wsw, "D%d" % n, "=C%d*%s" % (n, REF["custo_hora_dev"]), NEG, BRL0, fill=TOTAL)
cel(wsw, "E%d" % n, "Com a tela no WMS (A2) no lugar do app: troque o item A1 (160 h) pelo A2 (80 h).", wrap=True, fill=TOTAL)
faixa_total(wsw, n + 1, "TOTAL, Arquitetura B", "B")
cel(wsw, "C%d" % (n + 1), '=SUMIF(%s,"A e B",%s)+SUMIF(%s,"B",%s)' % (rng + rng), NEG, NUM, fill=TOTAL)
cel(wsw, "D%d" % (n + 1), "=C%d*%s" % (n + 1, REF["custo_hora_dev"]), NEG, BRL0, fill=TOTAL)
cel(wsw, "E%d" % (n + 1), None, fill=TOTAL)
cel(wsw, "A%d" % (n + 3), "Custo-hora de desenvolvimento calculado na aba Premissas (salário de mercado × encargos ÷ horas). "
    "É custo interno da equipe, não desembolso extra, a menos que o desenvolvimento seja contratado.", NOTA, borda=False)
SW_A, SW_B = n, n + 1

# =====================================================================
# 9b. POC (bancada + estações 04 e 27)
# =====================================================================
wq = wb.create_sheet("POC")
cabecalho(wq, 3, ["Item", "Base de cálculo", "Qtd", "Preço unit. (R$)", "Total (R$)"], [48, 58, 10, 16, 16])
cel(wq, "A1", "POC: bancada + estações 04 e 27 (Arquitetura A, coletor)", TIT, borda=False)
cel(wq, "A2", "Fase 1 valida o hardware e o software na bancada; fase 2 instala nas estações 04 (Medicamento) e 27 (Perfumaria). "
    "As estações da POC são as marcadas 'S' na aba Estações.", NOTA, borda=False)


def linha_poc(r, nome, base, qtd, preco, fmt=NUM):
    cel(wq, "A%d" % r, nome, wrap=True)
    cel(wq, "B%d" % r, base, wrap=True)
    cel(wq, "C%d" % r, qtd, AZUL if not str(qtd).startswith("=") else PRETO, fmt)
    cel(wq, "D%d" % r, "=" + preco, VERDE, BRL)
    cel(wq, "E%d" % r, "=C%d*D%d" % (r, r), PRETO, BRL)


r = 4
cel(wq, "A%d" % r, "Fase 1, bancada (1 coluna de alto giro, 1 zona de cada tipo)", SUB, borda=False)
r += 1
ini = r
bancada = [
    ("Concentrador 12 V completo", "gabinete, fonte, placa, alimentação", 1, UNIT["Concentrador 12 V"]),
    ("Nó de coluna", "1 prédio de alto giro", 1, UNIT["Nó de coluna"]),
    ("Módulos Q", "4 andares (0 a 3) do prédio", 4, UNIT["Módulo Q"]),
    ("Módulos E+Q", "1 zona dos andares 4-5 + 1 zona do baixo giro", 2, UNIT["Módulo E+Q"]),
    ("Sinal luminoso de baixo giro", "barra âmbar", 1, PRECO["P47"]),
    ("Cabo Cat5e (m)", "barramento e ligações da bancada", 20, PRECO["P12"]),
    ("Plugs RJ45", "", 20, PRECO["P11"]),
    ("Painel de MDF com longarina", "simula a prateleira", 1, PRECO["P48"]),
    ("Switch 5 portas", "rede da bancada", 1, PRECO["P49"]),
    ("Ferramentas de rede (crimpar e testar)", "compra única, servem para a fase 2", 1, PRECO["P41"]),
    ("Bancada de eletrônica (solda, multímetro)", "compra única, serve para a fase 2", 1, PRECO["P42"]),
]
for nome, base, q, pr in bancada:
    linha_poc(r, nome, base, q, pr)
    r += 1
cel(wq, "A%d" % r, "Frete das compras")
cel(wq, "B%d" % r, "% sobre os itens acima")
cel(wq, "C%d" % r, None)
cel(wq, "D%d" % r, "=" + REF["frete"], VERDE, PCT)
cel(wq, "E%d" % r, "=D%d*SUM(E%d:E%d)" % (r, ini, r - 1), PRETO, BRL)
r += 1
linha_poc(r, "Montagem dos módulos (h)", "conc. + nó + 4 Q + 2 E+Q",
          "=%s+%s+4*%s+2*%s" % (REF["h_conc"], REF["h_no"], REF["h_mod_q"], REF["h_mod_eq"]), REF["custo_hora"], NUM2)
r += 1
linha_poc(r, "Testes de bancada (h)", "acender, confirmar, falha, troca e cadastro de LED, 48 h de operação contínua", 16, REF["custo_hora"], NUM2)
r += 1
Q1 = r
faixa_total(wq, r, "TOTAL FASE 1, bancada", "BCD")
cel(wq, "E%d" % r, "=SUM(E%d:E%d)" % (ini, r - 1), NEG, BRL, fill=TOTAL)
r += 2
cel(wq, "A%d" % r, "Fase 2, instalação nas estações 04 e 27", SUB, borda=False)
r += 1
ini2 = r
cel(wq, "A%d" % r, "Materiais, montagem, instalação e treinamento das 2 estações")
cel(wq, "B%d" % r, "aba A-Interno, coluna piloto (estações com 'S')", wrap=True)
for col in "CD":
    cel(wq, "%s%d" % (col, r), None)
cel(wq, "E%d" % r, "='A-Interno'!F%d" % A_TOT, VERDE, BRL)
r += 1
cel(wq, "A%d" % r, "(−) Ferramentas já compradas na fase 1")
cel(wq, "B%d" % r, "kit de rede + bancada de eletrônica")
for col in "CD":
    cel(wq, "%s%d" % (col, r), None)
cel(wq, "E%d" % r, "=-('A-Interno'!F%d+'A-Interno'!F%d)*(1+%s)" % (LIN["ferr1"], LIN["ferr2"], REF["frete"]), PRETO, BRL)
r += 1
cel(wq, "A%d" % r, "(−) Módulos da bancada reaproveitados na estação")
cel(wq, "B%d" % r, "o concentrador, o nó, os 4 Q e os 2 E+Q da bancada vão para a estação 04", wrap=True)
for col in "CD":
    cel(wq, "%s%d" % (col, r), None)
cel(wq, "E%d" % r, "=-(%s+%s+4*%s+2*%s)" % (UNIT["Concentrador 12 V"], UNIT["Nó de coluna"], UNIT["Módulo Q"], UNIT["Módulo E+Q"]), PRETO, BRL)
r += 1
Q2 = r
faixa_total(wq, r, "TOTAL FASE 2, estações 04 e 27", "BCD")
cel(wq, "E%d" % r, "=SUM(E%d:E%d)" % (ini2, r - 1), NEG, BRL, fill=TOTAL)
r += 2
QH = r
faixa_total(wq, r, "HARDWARE DA POC (fase 1 + fase 2)", "BCD")
cel(wq, "E%d" % r, "=E%d+E%d" % (Q1, Q2), NEG, BRL, fill=TOTAL)
r += 1
QC = r
cel(wq, "A%d" % r, "Contingência")
cel(wq, "B%d" % r, "% sobre o hardware da POC")
cel(wq, "C%d" % r, None)
cel(wq, "D%d" % r, "=" + REF["contingencia"], VERDE, PCT)
cel(wq, "E%d" % r, "=D%d*E%d" % (r, QH), PRETO, BRL)
r += 1
QHC = r
faixa_total(wq, r, "HARDWARE DA POC COM CONTINGÊNCIA (desembolso)", "BCD")
cel(wq, "E%d" % r, "=E%d+E%d" % (QH, QC), NEG, BRL, fill=TOTAL)
r += 2
cel(wq, "A%d" % r, "Software para a POC (equipe interna, versão mínima)", SUB, borda=False)
r += 1
ini3 = r
swpoc = [
    ("Firmware do concentrador, dos nós e dos E+Q", "barramento, endereçamento, diagnóstico, cadastro", 160),
    ("Serviço PTL, versão mínima", "abrir e fechar caixa, acender, confirmar, sinal de baixo giro", 160),
    ("Integração Sybase (leitura dos itens da caixa)", "views de leitura e gravação de eventos", 60),
    ("App do coletor (A1), versão mínima", "leitura com gatilho e progresso da caixa", 100),
    ("Interface de controle, versão mínima", "falhas de LED e cadastro por endereço", 40),
    ("Testes da bancada e operação assistida", "ajustes durante a POC", 80),
]
for nome, base, h in swpoc:
    cel(wq, "A%d" % r, nome, wrap=True)
    cel(wq, "B%d" % r, base, wrap=True)
    cel(wq, "C%d" % r, h, AZUL, NUM, fill=AMARELO)
    cel(wq, "D%d" % r, "=" + REF["custo_hora_dev"], VERDE, BRL)
    cel(wq, "E%d" % r, "=C%d*D%d" % (r, r), PRETO, BRL)
    r += 1
QS = r
faixa_total(wq, r, "SOFTWARE DA POC (horas e custo interno)", "BD")
cel(wq, "C%d" % r, "=SUM(C%d:C%d)" % (ini3, r - 1), NEG, NUM, fill=TOTAL)
cel(wq, "E%d" % r, "=SUM(E%d:E%d)" % (ini3, r - 1), NEG, BRL, fill=TOTAL)
r += 2
cel(wq, "A%d" % r, "Prazo estimado: bancada 3 a 4 semanas; instalação nas 2 estações 1 semana; operação assistida 4 semanas.", NOTA, borda=False)
cel(wq, "A%d" % (r + 1), "O software da POC é a primeira parte do software do projeto (aba Software), não um custo a mais.", NOTA, borda=False)
wq.freeze_panes = "A4"

# =====================================================================
# 10. RESUMO
# =====================================================================
wr = wb.create_sheet("Resumo")
cabecalho(wr, 3, ["Cenário", "Descrição", "Piloto (R$)", "Operação completa (R$)", "Por estação, completo (R$)",
                  "Completo + contingência (R$)"], [42, 62, 17, 21, 21, 22])
cel(wr, "A1", "Resumo do orçamento, Pick-to-Light", TIT, borda=False)
cel(wr, "A2", "Piloto: estações marcadas 'S' na aba Estações (POC: 04 e 27). A operação completa já inclui o piloto. "
    "Hardware, instalação e treinamento; software nas linhas abaixo.", NOTA, borda=False)
cen = [
    ("A interno, só materiais (12 V)", "Peças, infraestrutura e frete; montagem própria", "'A-Interno'!F%d" % A_PECAS, "'A-Interno'!G%d" % A_PECAS),
    ("A interno, só materiais (variante 5 V)", "Sem reguladores, fonte 5 V, mais cabo de força", "'A-Interno'!F%d" % A_5V, "'A-Interno'!G%d" % A_5V),
    ("A interno, materiais + mão de obra", "Montagem, instalação, comissionamento e treinamento internos",
     "'A-Interno'!F%d" % A_TOT, "'A-Interno'!G%d" % A_TOT),
    ("A terceiros, T1 (PCB própria via EMS)", "Placas próprias fabricadas/montadas por terceiros + instalação",
     "'A-Terceiros'!F%d" % T1_TOT, "'A-Terceiros'!G%d" % T1_TOT),
    ("A terceiros, T2 (PTL comercial importado)", "Hardware PTL pronto + instalação",
     "'A-Terceiros'!F%d" % T2_TOT, "'A-Terceiros'!G%d" % T2_TOT),
    ("B1, RFID com pulseira", "Displays (A interno) + tags + pulseiras NFC", "'B-RFID'!F%d" % B1_TOT, "'B-RFID'!G%d" % B1_TOT),
    ("B2, RFID com leitor fixo na estação", "Displays (A interno) + tags + leitor fixo e painel (opção indicada na B)",
     "'B-RFID'!F%d" % B2_TOT, "'B-RFID'!G%d" % B2_TOT),
]
for i, (nome, desc, fp, fc) in enumerate(cen, 4):
    cel(wr, "A%d" % i, nome, NEG)
    cel(wr, "B%d" % i, desc, wrap=True)
    cel(wr, "C%d" % i, "=" + fp, VERDE, BRL0)
    cel(wr, "D%d" % i, "=" + fc, VERDE, BRL0)
    cel(wr, "E%d" % i, "=D%d/%s" % (i, tot("O", "C")), PRETO, BRL0)
    cel(wr, "F%d" % i, "=D%d*(1+%s)" % (i, REF["contingencia"]), PRETO, BRL0)
LB2 = 4 + [c[0] for c in cen].index("B2, RFID com leitor fixo na estação")
n = len(cen) + 4
cel(wr, "A%d" % (n + 1), "Desenvolvimento de software (interno)", SUB, borda=False)
cel(wr, "A%d" % (n + 2), "Arquitetura A (app próprio)", NEG)
cel(wr, "B%d" % (n + 2), "=\"Horas: \"&TEXT(Software!C%d,\"#,##0\")" % SW_A, PRETO)
cel(wr, "D%d" % (n + 2), "=Software!D%d" % SW_A, VERDE, BRL0)
cel(wr, "A%d" % (n + 3), "Arquitetura B", NEG)
cel(wr, "B%d" % (n + 3), "=\"Horas: \"&TEXT(Software!C%d,\"#,##0\")" % SW_B, PRETO)
cel(wr, "D%d" % (n + 3), "=Software!D%d" % SW_B, VERDE, BRL0)
LTOT = n + 5
LA = 4 + [c[0] for c in cen].index("A interno, materiais + mão de obra")
for k, (rot, desc, lh, ls) in enumerate([
        ("Investimento total, Arquitetura A (proposta)", "A interno com contingência + software da Arquitetura A", LA, n + 2),
        ("Investimento total, Arquitetura B2", "B2 com contingência + software da Arquitetura B", LB2, n + 3)]):
    rr = LTOT + k
    cel(wr, "A%d" % rr, rot, NEG, fill=TOTAL)
    cel(wr, "B%d" % rr, desc, fill=TOTAL, wrap=True)
    for col in "CDE":
        cel(wr, "%s%d" % (col, rr), None, fill=TOTAL)
    cel(wr, "F%d" % rr, "=F%d+D%d" % (lh, ls), NEG, BRL0, fill=TOTAL)
LTOT += 1
cel(wr, "A%d" % (LTOT + 2), "POC (bancada + estações 04 e 27)", NEG, fill=TOTAL)
cel(wr, "B%d" % (LTOT + 2), "Hardware com contingência (aba POC); software da POC é interno", fill=TOTAL, wrap=True)
cel(wr, "C%d" % (LTOT + 2), "=POC!E%d" % QHC, NEG, BRL0, fill=TOTAL)
for col in "DEF":
    cel(wr, "%s%d" % (col, LTOT + 2), None, fill=TOTAL)
LTOT += 2
cel(wr, "A%d" % (LTOT + 2), "Volume da operação (informado): ~3.000 volumes/dia e 250.000 a 270.000 produtos separados por dia.",
    NOTA, borda=False)
wr.freeze_panes = "A4"
wb.move_sheet("Resumo", offset=-(len(wb.sheetnames) - 2))

for w in wb.worksheets:
    w.sheet_view.showGridLines = False
wb.save(SAIDA)
print("ok", SAIDA, {"A_PECAS": A_PECAS, "A_TOT": A_TOT, "T1": T1_TOT, "T2": T2_TOT, "LTOT": LTOT})
